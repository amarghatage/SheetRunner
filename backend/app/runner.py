import html
import json
import os
import re
import time
import uuid
from datetime import datetime, timezone
from pathlib import Path

from selenium import webdriver
from selenium.common.exceptions import WebDriverException
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait, Select
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.common.action_chains import ActionChains

from .workbooks import load_plan

LOCATORS = {"ID": By.ID, "NAME": By.NAME, "XPATH": By.XPATH, "CSS": By.CSS_SELECTOR,
            "CSS SELECTOR": By.CSS_SELECTOR, "CLASS NAME": By.CLASS_NAME, "LINK TEXT": By.LINK_TEXT}


def _browser(name):
    name = name.lower()
    if name == "chrome": return webdriver.Chrome()
    if name == "firefox": return webdriver.Firefox()
    if name in ("edge", "msedge"): return webdriver.Edge()
    raise ValueError(f"Unsupported browser: {name}")


def _property_map(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb["CapturedObjectProperties"]
    rows = list(ws.iter_rows(values_only=True)); wb.close()
    heads = [str(x or "").strip().lower() for x in rows[0]]
    def ix(*names):
        for n in names:
            if n in heads: return heads.index(n)
    pi, ei = ix("page name", "page"), ix("element name", "element")
    si, vi = ix("locator strategy", "strategy", "locator type"), ix("locator value", "value", "element value")
    if None in (pi, ei, si, vi): raise ValueError("CapturedObjectProperties needs Page Name, Element Name, Locator Strategy, Locator Value columns")
    return {(str(r[pi]).strip(), str(r[ei]).strip()): (str(r[si]).strip().upper(), str(r[vi]).strip()) for r in rows[1:] if len(r) > vi and r[pi] and r[ei]}


def _run_action(driver, wait, step, locator, timeout):
    action = re.sub(r"[^a-z0-9]", "", step.action.lower())
    el = None
    if locator:
        strategy, value = locator
        if strategy not in LOCATORS: raise ValueError(f"Unknown locator strategy: {strategy}")
        el = wait.until(EC.presence_of_element_located((LOCATORS[strategy], value)))
    value = step.action_value
    if action in ("click", "clickbutton", "clicklink"): wait.until(EC.element_to_be_clickable((LOCATORS[locator[0]], locator[1]))).click()
    elif action in ("submit",): el.submit()
    elif action in ("entertext", "type", "sendkeys", "input"): el.clear(); el.send_keys(value)
    elif action in ("clear", "cleartext"): el.clear()
    elif action in ("verifytext", "gettext", "readtext"):
        actual = el.text
        if value and value.lower() not in actual.lower(): raise AssertionError(f"Expected text '{value}', got '{actual}'")
    elif action in ("verifyvalue", "getvalue"): 
        actual = el.get_attribute("value") or ""
        if value and actual != value: raise AssertionError(f"Expected value '{value}', got '{actual}'")
    elif action in ("verifytitle",):
        if value and value not in driver.title: raise AssertionError(f"Expected title '{value}', got '{driver.title}'")
    elif action in ("verifyurl",):
        if value and value not in driver.current_url: raise AssertionError(f"Expected URL '{value}', got '{driver.current_url}'")
    elif action in ("verifyvisible", "isvisible"):
        if not el.is_displayed(): raise AssertionError("Element is not visible")
    elif action in ("verifyenabled",):
        if not el.is_enabled(): raise AssertionError("Element is not enabled")
    elif action in ("verifychecked",):
        if not el.is_selected(): raise AssertionError("Element is not selected")
    elif action in ("acceptalert", "dismissalert", "verifyalerttext"):
        alert = wait.until(EC.alert_is_present())
        if action == "acceptalert": alert.accept()
        elif action == "dismissalert": alert.dismiss()
        elif value not in alert.text: raise AssertionError(f"Alert text mismatch: {alert.text}")
    elif action in ("selectbytext", "selectbyvalue", "selectbyindex"):
        sel = Select(el)
        {"selectbytext": sel.select_by_visible_text, "selectbyvalue": sel.select_by_value,
         "selectbyindex": lambda v: sel.select_by_index(int(v))}[action](value)
    elif action == "navigate": driver.get(value)
    elif action == "back": driver.back()
    elif action == "forward": driver.forward()
    elif action == "refresh": driver.refresh()
    elif action == "maximize": driver.maximize_window()
    elif action == "wait": time.sleep(float(value or 1))
    elif action in ("scrolltoelement", "scroll"):
        driver.execute_script("arguments[0].scrollIntoView({block:'center'});", el)
    elif action == "hover": ActionChains(driver).move_to_element(el).perform()
    elif action == "doubleclick": ActionChains(driver).double_click(el).perform()
    elif action == "rightclick": ActionChains(driver).context_click(el).perform()
    elif action in ("switchtoframe", "frame"): driver.switch_to.frame(el)
    elif action in ("switchtoparentframe", "parentframe"): driver.switch_to.parent_frame()
    elif action in ("switchtodefaultcontent", "defaultcontent"): driver.switch_to.default_content()
    elif action in ("switchtowindow", "window"): driver.switch_to.window(driver.window_handles[-1])
    elif action in ("close",): driver.close()
    else: raise ValueError(f"Unknown action: {step.action}")


def execute_run(run_id, config, jobs):
    started = datetime.now(timezone.utc).isoformat()
    report_dir, shot_dir = Path(config.output_dir), Path(config.screenshot_dir)
    report_dir.mkdir(parents=True, exist_ok=True); shot_dir.mkdir(parents=True, exist_ok=True)
    records = []; driver = None; disabled = []
    try:
        steps, disabled = load_plan(config.suite_path, config.cases_path)
        for case in disabled:
            records.append({"case": case, "step": "", "action": "", "status": "skipped", "details": "Execution is NO in the suite workbook."})
        properties = _property_map(config.cases_path)
        cases = {}
        for step in steps: cases.setdefault(step.case, []).append(step)
        driver = _browser(config.browser)
        driver.set_window_size(1440, 900)
        try:
            driver.get(config.target_url)
        except WebDriverException as exc:
            message = str(exc)
            if "ERR_NAME_NOT_RESOLVED" in message:
                raise ValueError(f"Chrome could not resolve the target host in '{config.target_url}'. Check the URL spelling and confirm the site is reachable from this computer.") from exc
            raise ValueError(f"Could not open target URL '{config.target_url}': {message.splitlines()[0]}") from exc
        for case, case_steps in cases.items():
            for step in case_steps:
                now = datetime.now(timezone.utc).isoformat(); t = time.monotonic()
                record = {"case": case, "step": step.step_id or step.action, "action": step.action, "started_at": now}
                try:
                    locator = properties.get((step.page, step.element)) if step.page and step.element else None
                    if step.page and step.element and not locator: raise ValueError(f"No captured locator for {step.page}.{step.element}")
                    _run_action(driver, WebDriverWait(driver, config.timeout_seconds), step, locator, config.timeout_seconds)
                    record["status"] = "passed"
                except Exception as exc:
                    record["status"] = "failed"; record["error"] = str(exc)
                    shot = shot_dir / f"{run_id}_{len(records)+1}.png"
                    try:
                        driver.save_screenshot(str(shot)); record["screenshot"] = str(shot)
                    except Exception: pass
                    if step.on_failure.upper() not in ("CONTINUE", "SKIP", "IGNORE"):
                        records.append(record); record["duration_seconds"] = round(time.monotonic()-t, 3)
                        raise
                record["duration_seconds"] = round(time.monotonic()-t, 3); records.append(record)
        statuses = [r["status"] for r in records]
        overall = "failed" if "failed" in statuses else ("passed" if "passed" in statuses else "skipped")
        result = {"run_id": run_id, "status": overall, "started_at": started, "finished_at": datetime.now(timezone.utc).isoformat(), "disabled_cases": disabled, "steps": records}
    except Exception as exc:
        result = {"run_id": run_id, "status": "failed", "started_at": started, "finished_at": datetime.now(timezone.utc).isoformat(), "disabled_cases": disabled, "error": str(exc), "steps": records}
    finally:
        if driver:
            try: driver.quit()
            except Exception: pass
    (report_dir / f"{run_id}.json").write_text(json.dumps(result, indent=2), encoding="utf-8")
    rows = "".join(f"<tr><td>{html.escape(s.get('case',''))}</td><td>{html.escape(s.get('step',''))}</td><td>{html.escape(s.get('action',''))}</td><td>{html.escape(s.get('status',''))}</td><td>{html.escape(s.get('error') or s.get('details',''))}</td></tr>" for s in result.get("steps", []))
    if not rows:
        rows = "<tr><td colspan='5'>No test cases were run.</td></tr>"
    error_box = f"<p style='color:#b91c1c'><strong>Run error:</strong> {html.escape(result['error'])}</p>" if result.get("error") else ""
    summary = f"<p>Passed: {sum(s.get('status')=='passed' for s in records)} · Failed: {sum(s.get('status')=='failed' for s in records)} · Skipped: {sum(s.get('status')=='skipped' for s in records)}</p>"
    page = f"<!doctype html><html><head><meta charset='utf-8'><title>Run {html.escape(run_id)}</title><style>body{{font:15px system-ui;margin:32px;color:#182230}}table{{border-collapse:collapse;width:100%}}th,td{{border:1px solid #d5dbe5;padding:8px;text-align:left}}th{{background:#f1f4f8}}</style></head><body><h1>Run {html.escape(run_id)}: {html.escape(result['status'])}</h1>{summary}{error_box}<table><tr><th>Case</th><th>Step</th><th>Action</th><th>Status</th><th>Details</th></tr>{rows}</table></body></html>"
    (report_dir / f"{run_id}.html").write_text(page, encoding="utf-8")
    junit_cases = "".join(f"<testcase classname='{html.escape(s.get('case',''))}' name='{html.escape(s.get('step',''))}' time='{s.get('duration_seconds',0)}'>" + (f"<failure message='{html.escape(s.get('error',''))}'/>" if s.get("status") == "failed" else "") + "</testcase>" for s in records)
    (report_dir / f"{run_id}.xml").write_text(f"<testsuite tests='{len(records)}' failures='{sum(s.get('status')=='failed' for s in records)}'>{junit_cases}</testsuite>", encoding="utf-8")
    jobs[run_id].update({"status": result["status"], "result": result, "report": str(report_dir / f"{run_id}.html")})
