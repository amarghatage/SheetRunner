from pathlib import Path
from openpyxl import Workbook

out = Path(__file__).resolve().parents[2] / "sample_workbooks"
out.mkdir(exist_ok=True)

suite = Workbook(); ws = suite.active; ws.title = "Suite"
ws.append(["Test Case", "Execution"]); ws.append(["Open homepage", "YES"]); ws.append(["Disabled example", "NO"])
suite.save(out / "suite.xlsx")

cases = Workbook(); ws = cases.active; ws.title = "TestCaseSheet"
ws.append(["Test Case", "Step ID", "Page Name", "Locator Strategy", "Element Name", "Action", "Action Value", "On-Failure Behavior", "Test-Data Reference"])
ws.append(["Open homepage", "1", "", "", "", "verifyTitle", "Example Domain", "FAIL", ""])
props = cases.create_sheet("CapturedObjectProperties")
props.append(["Page Name", "Element Name", "Locator Strategy", "Locator Value"])
props.append(["Home", "Search", "NAME", "q"])
cases.create_sheet("InputData").append(["Query"])
cases.save(out / "test-cases.xlsx")
print(f"Sample workbooks created in {out}")
