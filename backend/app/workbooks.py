from dataclasses import dataclass
from pathlib import Path
from openpyxl import load_workbook


@dataclass
class Step:
    case: str
    step_id: str
    page: str
    strategy: str
    element: str
    action: str
    action_value: str = ""
    on_failure: str = "FAIL"
    data_ref: str = ""


def _text(value):
    return "" if value is None else str(value).strip()


def _rows(path, sheet=None):
    book = load_workbook(path, data_only=True, read_only=True)
    ws = book[sheet] if sheet and sheet in book.sheetnames else book.active
    rows = list(ws.iter_rows(values_only=True))
    book.close()
    if not rows:
        raise ValueError(f"Workbook '{Path(path).name}' is empty")
    headers = [_text(v).lower() for v in rows[0]]
    return headers, rows[1:]


def _col(headers, *aliases):
    for alias in aliases:
        if alias.lower() in headers:
            return headers.index(alias.lower())
    return None


def validate_workbooks(suite_path, cases_path):
    issues = []
    try:
        sh, sr = _rows(suite_path)
        case_i = _col(sh, "test case", "testcase", "test case name", "testcasename", "case")
        exec_i = _col(sh, "execution", "execute", "status", "run")
        if case_i is None or exec_i is None:
            issues.append("Suite sheet needs Test Case and Execution columns.")
        ch, cr = _rows(cases_path, "TestCaseSheet")
        required = {
            "case": _col(ch, "test case", "testcase", "test case name", "testcasename"),
            "id": _col(ch, "step id", "step", "step name"),
            "page": _col(ch, "page name", "page"),
            "strategy": _col(ch, "locator strategy", "strategy", "locator type"),
            "element": _col(ch, "element name", "element", "locator value"),
            "action": _col(ch, "action", "action keyword"),
        }
        missing = [key for key, idx in required.items() if idx is None]
        if missing:
            issues.append("TestCaseSheet is missing required columns: " + ", ".join(missing))
        if "CapturedObjectProperties" not in load_workbook(cases_path, read_only=True).sheetnames:
            issues.append("Test-case workbook needs a CapturedObjectProperties sheet.")
        if not issues:
            suite_cases = {
                _text(row[case_i])
                for row in sr
                if len(row) > max(case_i, exec_i)
                and _text(row[case_i])
                and _text(row[exec_i]).upper() == "YES"
            }
            defined = {_text(row[required["case"]]) for row in cr if len(row) > required["case"] and _text(row[required["case"]])}
            unknown = sorted(suite_cases - defined)
            if unknown:
                issues.append("Suite cases missing from TestCaseSheet: " + ", ".join(unknown))
    except Exception as e:
        issues.append(str(e))
    return {"valid": not issues, "issues": issues}


def load_plan(suite_path, cases_path):
    errors = validate_workbooks(suite_path, cases_path)
    if not errors["valid"]:
        raise ValueError("; ".join(errors["issues"]))
    sh, sr = _rows(suite_path)
    si = _col(sh, "test case", "testcase", "test case name", "testcasename", "case")
    ei = _col(sh, "execution", "execute", "status", "run")
    selected = []
    disabled = []
    for row in sr:
        if len(row) <= max(si, ei):
            continue
        name, flag = _text(row[si]), _text(row[ei]).upper()
        if name and flag == "YES":
            selected.append(name)
        elif name:
            disabled.append(name)
    headers, rows = _rows(cases_path, "TestCaseSheet")
    indexes = {key: _col(headers, *aliases) for key, aliases in {
        "case": ("test case", "testcase", "test case name", "testcasename"),
        "id": ("step id", "step", "step name"), "page": ("page name", "page"),
        "strategy": ("locator strategy", "strategy", "locator type"),
        "element": ("element name", "element", "locator value"),
        "action": ("action", "action keyword"),
        "value": ("action value", "input value", "expected value"),
        "failure": ("on-failure behavior", "on failure", "failure behavior"),
        "data": ("test-data reference", "data reference", "test data"),
    }.items()}
    output = []
    current = ""
    for row in rows:
        def v(key):
            i = indexes[key]
            return _text(row[i]) if i is not None and i < len(row) else ""
        current = v("case") or current
        if current in selected and v("action"):
            output.append(Step(current, v("id"), v("page"), v("strategy").upper(), v("element"), v("action"), v("value"), v("failure") or "FAIL", v("data")))
    return output, disabled

